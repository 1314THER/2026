#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""斜率截距比 -> 收缩判据：本文件夹全部数字的唯一来源。

论点
----
题目给出两条体密度经验式  rho = a + b*C。由干基含水率定义可写出单位干物质
的比体积

        v(C) = V / m_d = (1 + C) / (a + b C) = (1/a) * (1 + C) / (1 + r C),
        r  = b / a        （斜率截距比，无量纲）

于是：
  * 收缩规律的**形状**只由 r 决定，a 只提供尺度（rho(0) = a 为绝干表观密度）；
  * r = 1  <=>  v 为常数  <=>  无收缩（等价于 rho_d = m_d / V 为常数）；
  * r < 1  <=>  随干燥收缩；r 越小收缩越强。

这样一来，"要不要考虑收缩"就不再是一个定性判断，而是一个可以直接从题面
系数读出的定量判据。本脚本算出该判据的全部数值，并把它与附件 2 的实测
半径收缩做交叉验证，供论文中"问题四之后的思考"一节直接引用。

口径约定（重要）
----------------
问题四的几何是**长度不变、只缩半径**（附件 2 只给半径），故

        V / V0 = (R / R0)^2          （径向等长，本文采用）
        V / V0 = (R / R0)^3          （各向同性，仅作对照）

两种口径给出的线收缩比不同，全文必须只用一种；本脚本两种都算，正文只用径向。

运行
----
    PY=./.venv/bin/python
    $PY compute_shrinkage.py                 # 写 输出/ 下的 json 与 csv
    $PY compute_shrinkage.py --print         # 同时打印到终端
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[1]   # 包根目录（脚本放在 补充分析/斜率截距与收缩/ 下）
OUTDIR = HERE / "输出"

# ------------------------------------------------------------------ 常数
C0 = 2.55            # 初始干基含水率 (kg/kg)
RHO_W = 1000.0       # 水的真密度 (kg/m^3)
RHO_S = 1400.0       # 干物质真密度 (kg/m^3)，仅用于三相检验的对照
R0_CM = 2.000        # 初始半径 (cm)

CASES = {
    "附录3": {"a": 650.0, "b": 128.0},
    "附录4": {"a": 760.0, "b": 90.0},
}

# 线收缩门槛（用于"多大含水率区间内可以忽略收缩"这张表）
EPS_LIST = (0.01, 0.02, 0.03, 0.05, 0.10, 0.15, 0.20)


# ------------------------------------------------------------------ 解析式
def r_ratio(a: float, b: float) -> float:
    """斜率截距比 r = b / a（无量纲）。"""
    return b / a


def v_specific(C, a: float, b: float):
    """单位干物质比体积 v(C) = (1+C)/(a+bC)，单位 m^3/kg 干料。"""
    return (1.0 + np.asarray(C, dtype=float)) / (a + b * np.asarray(C, dtype=float))


def v_ratio(C, r: float, c0: float = C0):
    """V(C)/V(C0) = (1+C)(1+r*C0) / [(1+C0)(1+r*C)]。"""
    C = np.asarray(C, dtype=float)
    return (1.0 + C) * (1.0 + r * c0) / ((1.0 + c0) * (1.0 + r * C))


def v_ratio_dry(r: float, c0: float = C0) -> float:
    """绝干外推 V(0)/V(C0) = (1 + r*C0) / (1 + C0)。"""
    return (1.0 + r * c0) / (1.0 + c0)


def line_shrink(C, r: float, geom: str = "radial", c0: float = C0):
    """线收缩比 1 - L/L0。geom='radial' 用平方根（等长），'iso' 用立方根。"""
    q = v_ratio(C, r, c0)
    root = np.sqrt(q) if geom == "radial" else np.cbrt(q)
    return 1.0 - root


def beta_endpoint(a: float, b: float, c0: float = C0, rho_w: float = RHO_W) -> float:
    """端点反演的收缩系数 beta = rho_w (a-b) / [a (a + b C0)]。"""
    return rho_w * (a - b) / (a * (a + b * c0))


def b_band(a: float, c0: float = C0, rho_w: float = RHO_W):
    """beta in [0,1] 等价于斜率 b 落在 [b_min, a]；返回 (b_min, a)。"""
    b_min = a * (rho_w - a) / (rho_w + a * c0)
    return b_min, a


def C_at_line_shrink(eps: float, r: float, c0: float = C0) -> float:
    """解 1 - sqrt(V(C)/V(C0)) = eps 的含水率 C（径向口径，闭式）。"""
    k = 1.0 + r * c0
    m = (1.0 - eps) ** 2 * (1.0 + c0)
    if abs(k - m * r) < 1e-15:
        return float("nan")
    return (m - k) / (k - m * r)


def naive_no_shrink_rho(a: float, c0: float = C0) -> float:
    """若强行要求斜率=截距（无收缩），C0 处的体密度 rho = a(1+C0)。"""
    return a * (1.0 + c0)


# ------------------------------------------------------------------ 数据读入
def read_radius_cm():
    """读附件 2 的半径序列（xlsx），返回 (t_s, R_cm) 两个 ndarray。"""
    from openpyxl import load_workbook

    xlsx = PROJECT / "附件" / "附件2-药材半径.xlsx"
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    ts, rs = [], []
    for row in ws.iter_rows(values_only=True):
        if row is None or len(row) < 2:
            continue
        try:
            t = float(row[0]); r = float(row[1])
        except (TypeError, ValueError):
            continue                  # 表头
        ts.append(t); rs.append(r)
    wb.close()
    return np.asarray(ts), np.asarray(rs)


def read_result4():
    """读问题四的结果场（npz），返回 times/conc21/radius/valid21。"""
    npz = PROJECT / "输出" / "result4_data.npz"
    d = np.load(npz, allow_pickle=True)
    return d["times"], d["conc21"], d["radius"], d["valid21"]


def _cbar(conc_row, r_cm_arr, R_cm_val):
    """按体积权重（~r）求半径 R_cm_val 以内的平均干基含水率。"""
    idx = [k for k in range(len(r_cm_arr)) if r_cm_arr[k] <= R_cm_val + 1e-9]
    if not idx:
        return float("nan")
    w = np.maximum(r_cm_arr[idx], 1e-9)
    return float(np.sum(np.asarray(conc_row)[idx] * w) / np.sum(w))


def build_echo() -> dict:
    """问题一/二/三的报数点 -> 由 r 反推的径向线收缩（用于三问呼应表）。"""
    d1 = np.load(PROJECT / "输出" / "result1_data.npz", allow_pickle=True)
    x_cm = np.arange(21) * 0.1
    echo = {}

    c = d1["conc_out"][1800]                      # 问题一：t = 1800 s
    echo["问题一"] = {
        "时刻": "1800 s",
        "中心C": float(c[0]), "表面C": float(c[-1]), "平均C": _cbar(c, x_cm, R0_CM),
    }

    d23 = np.load(PROJECT / "输出" / "result23_data.npz", allow_pickle=True)
    dt, t23, c23 = float(d23["dt"][0]), d23["times"], d23["conc21"]
    i3h = int(round(10800.0 / dt))
    c = c23[i3h]                                  # 问题二：t = 3 h
    echo["问题二"] = {
        "时刻": "3 h", "中心C": float(c[0]), "表面C": float(c[-1]),
        "平均C": _cbar(c, x_cm, R0_CM),
    }
    c = c23[-1]                                   # 问题三：烘干判据达成
    echo["问题三"] = {
        "时刻": f"{float(d23['dry_time'][0]) / 3600:.2f} h",
        "中心C": float(c[0]), "表面C": float(c[-1]), "平均C": _cbar(c, x_cm, R0_CM),
    }

    for name, rec in echo.items():
        for case, prm in CASES.items():
            r = r_ratio(prm["a"], prm["b"])
            rec[f"{case}_线收缩_表面"] = float(line_shrink(rec["表面C"], r, "radial"))
            rec[f"{case}_线收缩_平均"] = float(line_shrink(rec["平均C"], r, "radial"))
    return echo


# ------------------------------------------------------------------ 主计算
def build_report() -> dict:
    rep: dict = {
        "说明": "斜率截距比 r=b/a 作为是否考虑收缩的判据；口径=径向等长",
        "常数": {"C0": C0, "rho_w": RHO_W, "rho_s": RHO_S, "R0_cm": R0_CM},
        "附录": {},
        "附件2": {},
        "门槛表": {},
    }

    for name, prm in CASES.items():
        a, b = prm["a"], prm["b"]
        r = r_ratio(a, b)
        bmin, bmax = b_band(a)
        rec = {
            "a": a, "b": b, "r": r,
            "rho_C0": a + b * C0,
            "V_dry_over_V_C0": v_ratio_dry(r),
            "R_dry_cm_radial": R0_CM * math.sqrt(v_ratio_dry(r)),
            "R_dry_cm_iso": R0_CM * v_ratio_dry(r) ** (1.0 / 3.0),
            "beta_endpoint": beta_endpoint(a, b),
            "b_band_lo": bmin, "b_band_hi": bmax,
            "b_position_in_band": (b - bmin) / (bmax - bmin),
            "rho_if_no_shrink": naive_no_shrink_rho(a),
            "线收缩": {},
        }
        for C in (2.00, 1.51, 1.00, 0.50, 0.15):
            rec["线收缩"][f"C={C:.2f}"] = {
                "radial": float(line_shrink(C, r, "radial")),
                "iso": float(line_shrink(C, r, "iso")),
            }
        rep["附录"][name] = rec

    # ---- 附件 2 实测收缩
    t_s, R_cm = read_radius_cm()
    area = (R_cm[-1] / R_cm[0]) ** 2
    iso = (R_cm[-1] / R_cm[0]) ** 3
    at2 = {
        "n_points": int(len(t_s)),
        "R0_cm": float(R_cm[0]), "R_end_cm": float(R_cm[-1]),
        "V_end_over_V0_radial": float(area),
        "V_end_over_V0_iso": float(iso),
        "反标定beta": {},
        "绝干半径对照": {},
        "收缩完成度": {},
    }
    # 收缩完成度：R 的收缩量达到全程收缩量的 50% / 90% / 99% 所需的时间
    done = (R_cm[0] - R_cm) / (R_cm[0] - R_cm[-1])
    for level in (0.50, 0.90, 0.99):
        idx = int(np.argmax(done >= level))
        at2["收缩完成度"][f"{int(level*100)}%"] = {
            "t_h": float(t_s[idx] / 3600.0), "R_cm": float(R_cm[idx]),
        }
    for name, prm in CASES.items():
        a = prm["a"]
        at2["反标定beta"][name] = (1.0 / area - 1.0) * RHO_W / (a * C0)
        at2["绝干半径对照"][name] = {
            "R_dry_cm": R0_CM * math.sqrt(v_ratio_dry(r_ratio(prm["a"], prm["b"]))),
            "偏差_pct": 100.0 * (R0_CM * math.sqrt(v_ratio_dry(r_ratio(prm["a"], prm["b"])))
                                 - R_cm[-1]) / R_cm[-1],
        }
    rep["附件2"] = at2

    # ---- 线收缩门槛：多小的含水率区间可以忽略收缩
    thr = {}
    for name, prm in CASES.items():
        r = r_ratio(prm["a"], prm["b"])
        thr[name] = {f"{int(e*100)}%": float(C_at_line_shrink(e, r)) for e in EPS_LIST}
    rep["门槛表"] = thr

    # ---- 路径对照：同一含水率下 实测比体积 vs 由 rho 律反推的比体积
    t4, C4, R4, valid4 = read_result4()
    x_cm = np.arange(21) * 0.1                      # 输出列 = 固定物理距离(cm)
    r4_arr = np.asarray(R4, dtype=float) * 100.0    # m -> cm
    md0 = float((CASES["附录4"]["a"] + CASES["附录4"]["b"] * C0) * math.pi * (R0_CM * 1e-2) ** 2 * 0.25 / (1 + C0))
    path = []
    for i in range(0, len(t4), 3600):
        Rc = r4_arr[i]
        idx = [k for k in range(21) if x_cm[k] <= Rc + 1e-9]
        if len(idx) < 2:
            continue
        w = np.maximum(x_cm[idx], 1e-9)             # 体积权重 ~ r
        Cbar = float(np.sum(C4[i][idx] * w) / np.sum(w))
        v_meas = math.pi * (Rc * 1e-2) ** 2 * 0.25 / md0
        v_rho4 = float(v_specific(Cbar, CASES["附录4"]["a"], CASES["附录4"]["b"]))
        path.append({
            "t_h": float(t4[i] / 3600.0), "R_cm": float(Rc), "C_bar": Cbar,
            "v_meas_e3": v_meas * 1e3, "v_rho4_e3": v_rho4 * 1e3,
            "dev_pct": 100.0 * (v_meas - v_rho4) / v_rho4,
        })
    rep["路径对照"] = path
    devs = [p["dev_pct"] for p in path if p["t_h"] > 0.0]
    rep["路径偏差范围_pct"] = {"min": min(devs), "max": max(devs)}

    # ---- 三问呼应表：问题一/二/三的报数点对应的隐含径向线收缩
    rep["呼应表"] = build_echo()

    return rep


# ------------------------------------------------------------------ 输出
def write_outputs(rep: dict) -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    (OUTDIR / "斜率截距与收缩.json").write_text(
        json.dumps(rep, ensure_ascii=False, indent=2), encoding="utf-8")

    # 表 1：判据表（论文表 A）
    with (OUTDIR / "表A-判据.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["项目", "附录3", "附录4", "附件2实测"])
        rows = [
            ("截距 a = ρ(0) / (kg/m³)", "{:.0f}", "a"),
            ("斜率 b / (kg/m³)", "{:.0f}", "b"),
            ("斜率截距比 r = b/a", "{:.4f}", "r"),
            ("绝干体积比 V(0)/V(C₀)", "{:.4f}", "V_dry_over_V_C0"),
            ("绝干半径（径向口径）/ cm", "{:.3f}", "R_dry_cm_radial"),
            ("收缩系数 β（端点反演）", "{:.3f}", "beta_endpoint"),
        ]
        for label, fmt, key in rows:
            w.writerow([label] + [fmt.format(rep["附录"][k][key]) for k in ("附录3", "附录4")] + [""])
        w.writerow(["斜率允许带 [b_min, a]",
                    f"[{rep['附录']['附录3']['b_band_lo']:.1f}, {rep['附录']['附录3']['b_band_hi']:.0f}]",
                    f"[{rep['附录']['附录4']['b_band_lo']:.1f}, {rep['附录']['附录4']['b_band_hi']:.0f}]", ""])
        w.writerow(["实测终半径 / cm", "", "",
                    f"{rep['附件2']['R_end_cm']:.3f}"])
        w.writerow(["实测径向面积比 V_end/V0", "", "",
                    f"{rep['附件2']['V_end_over_V0_radial']:.4f}"])

    # 表 2：线收缩门槛表
    with (OUTDIR / "表B-门槛.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["线收缩达到", "附录3 对应 C", "附录4 对应 C"])
        for e in EPS_LIST:
            k = f"{int(e*100)}%"
            w.writerow([k, f"{rep['门槛表']['附录3'][k]:.3f}", f"{rep['门槛表']['附录4'][k]:.3f}"])

    # 表 3：路径对照
    with (OUTDIR / "表C-路径对照.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["t / h", "R / cm", "体积平均 C", "实测比体积 /(10⁻³m³/kg)",
                    "ρ律比体积 /(10⁻³m³/kg)", "相对偏差 / %"])
        for p in rep["路径对照"]:
            w.writerow([f"{p['t_h']:.1f}", f"{p['R_cm']:.3f}", f"{p['C_bar']:.3f}",
                        f"{p['v_meas_e3']:.4f}", f"{p['v_rho4_e3']:.4f}", f"{p['dev_pct']:+.1f}"])


def print_report(rep: dict) -> None:
    for name, rec in rep["附录"].items():
        print(f"{name}: a={rec['a']:.0f} b={rec['b']:.0f}  r={rec['r']:.4f}  "
              f"V(0)/V(C0)={rec['V_dry_over_V_C0']:.4f}  "
              f"R_dry(径向)={rec['R_dry_cm_radial']:.3f}cm  beta={rec['beta_endpoint']:.3f}  "
              f"斜率带=[{rec['b_band_lo']:.1f},{rec['b_band_hi']:.0f}]")
    a2 = rep["附件2"]
    print(f"附件2: n={a2['n_points']}  R {a2['R0_cm']:.3f}->{a2['R_end_cm']:.3f} cm  "
          f"面积比={a2['V_end_over_V0_radial']:.4f}  "
          f"反标定beta(附4)={a2['反标定beta']['附录4']:.3f}")
    for k, v in rep["门槛表"]["附录4"].items():
        print(f"  线收缩 {k:>4} 对应 C = {v:.3f}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--print", action="store_true", dest="do_print")
    args = ap.parse_args()
    rep = build_report()
    write_outputs(rep)
    if args.do_print:
        print_report(rep)
    print(f"已写出 -> {OUTDIR}")


if __name__ == "__main__":
    main()
