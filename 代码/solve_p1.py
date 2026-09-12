#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2026 高教社杯全国大学生数学建模竞赛 A 题 问题一求解.

控制方程（圆柱坐标，一维径向轴对称）:
    热量:  rho * cp * dT/dt = (1/r) * d/dr ( k * r * dT/dr )
    水分:  dC/dt             = (1/r) * d/dr ( D(C) * r * dC/dr )

定解条件:
    T(r,0) = 28 C,  C(r,0) = 2.55 kg/kg
    r = 0 :  dT/dr = 0,  dC/dr = 0                    (对称)
    r = R :  -k dT/dr = h  * (T_s - T_inf(t))         (对流换热)
             -D dC/dr = hm * (C_s - C_inf(t))         (对流传质)

T_inf(t)、C_inf(t) 由附件 1 线性插值给出。物性取自附录 2，其中
    D(C) = 7e-9 * exp(-0.89 / C)   [m^2/s]

建模选择（详见 建模/问题一-物理模型与公式.md）:
1. 传质边界采用干基密度口径，rho_d 两侧约去，故写 -D dC/dr = hm (C_s - C_inf)。
2. 忽略汽化潜热。若计入，按同一密度口径会把表面压到湿球温度(34.7 C)以下，物理不可能。
3. 水分方程取简单扩散形式，即 D 的经验公式视为相对总密度定义。

数值方法:
节点中心有限体积（离散格式守恒）+ 后向 Euler（无条件稳定），
D(C) 的非线性用 Picard 迭代。默认细网格 N=400（h = 0.005 cm），
题目要求的输出节点 0, 0.1, ..., 2.0 cm 恰为网格子集，无需插值。

运行环境（系统 python3 缺 openpyxl，请用打包运行时）:
    PY=./.venv/bin/python
    $PY 代码/solve_p1.py --save 输出/result1_data.npz
"""

from __future__ import annotations
import argparse
from pathlib import Path
import numpy as np
import openpyxl


PROJECT = Path(__file__).resolve().parent.parent
ATTACHMENT1 = PROJECT / "附件" / "附件1-烘房温度与水分浓度.xlsx"

RHO = 820.0          # 体积密度      kg/m^3
CP = 2600.0          # 比热容        J/(kg K)
K_COND = 0.36        # 热传导系数    W/(m K)
H_HEAT = 25.0        # 对流换热系数  W/(m^2 K)
H_MASS = 8.0e-7      # 对流传质系数  m/s

RADIUS = 0.02        # 半径          m
LENGTH = 0.25        # 长度          m

T_INIT = 28.0        # 药材温度初值   C
C_INIT = 2.55        # 水分浓度初值   kg/kg

T_END = 1800.0       # 模拟时长      s
DT = 1.0             # 时间步长      s

N_FINE = 400         # 细网格单元数，为 N_OUT 的整数倍，避免插值
N_OUT = 20           # 输出网格单元数（0.1 cm 间隔，21 个节点）

ALPHA = K_COND / (RHO * CP)

REPORT_TIMES = [100, 300, 600, 900, 1200, 1500, 1800]
REPORT_POS_CM = [0.0, 0.5, 1.0, 1.5, 2.0]


def load_air_conditions(path=ATTACHMENT1):
    #读取附件1，得到边界条件
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = [
        row
        for row in wb.active.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    t = np.array([row[0] for row in rows], dtype=float)
    temp = np.array([row[1] for row in rows], dtype=float)
    humid = np.array([row[2] for row in rows], dtype=float)
    return t, temp, humid


def diffusion_coefficient(c):
    #附录 2 的水分浓度扩散系数经验公式 [m^2/s]
    return 7.0e-9 * np.exp(-0.89 / np.maximum(c, 1e-12))    #防止C=0导致除零


def cell_weights(n_fine, hx):
    #用于守恒检验
    """有限体积控制体的径向权重 w_i，满足 ∫f dV = 2*pi*L * Σ w_i f_i。

    内部节点:  w_i = r_i * hx
    中心节点 (i=0): 控制体为 [0, hx/2]，V = pi*hx^2*L/4 -> w_0 = hx^2/8
    表面节点 (i=N): 控制体为 [R-hx/2, R]，V = pi*hx*L*(R-hx/4) -> w_N = hx*(R-hx/4)/2

    特别地，r_0 = 0，直接用 r_i*hx 会把中心控制体权重记成 0，导致守恒检验出现虚假残差。
    """
    r = np.arange(n_fine + 1) * hx
    w = r * hx
    w[0] = hx * hx / 8.0
    w[n_fine] = hx * (RADIUS - hx / 4.0) / 2.0
    return w


def thomas(lower, diag, upper, rhs):
    #追赶法求解三对角方程组
    n = len(diag)
    cp = np.empty(n)
    dp = np.empty(n)
    cp[0] = upper[0] / diag[0]
    dp[0] = rhs[0] / diag[0]
    for i in range(1, n):
        m = diag[i] - lower[i] * cp[i - 1]
        cp[i] = upper[i] / m
        dp[i] = (rhs[i] - lower[i] * dp[i - 1]) / m
    x = np.empty(n)
    x[-1] = dp[-1]
    for i in range(n - 2, -1, -1):
        x[i] = dp[i] - cp[i] * x[i + 1]
    return x


def step_temperature(temp, t_inf, dt, hx, r, rf):
    '''1.用后向 Euler 进行时间离散，无条件稳定
    2.得到线性三对角方程组，不需要迭代，直接用 thomas 求解
    '''
    n = len(r) - 1
    a_t = K_COND * dt / (RHO * CP * hx * hx)
    g_t = dt / (RHO * CP * hx * (RADIUS - hx / 4.0))
    p_t = 2.0 * g_t * RADIUS * H_HEAT
    q_t = 2.0 * g_t * (RADIUS - hx / 2.0) * K_COND / hx

    lower = np.zeros(n + 1)
    diag = np.zeros(n + 1)
    upper = np.zeros(n + 1)
    rhs = np.zeros(n + 1)

    diag[0] = 1.0 + 4.0 * a_t
    upper[0] = -4.0 * a_t
    rhs[0] = temp[0]
    for i in range(1, n):
        lower[i] = -a_t * rf[i - 1]
        upper[i] = -a_t * rf[i]
        diag[i] = r[i] + a_t * (rf[i - 1] + rf[i])
        rhs[i] = r[i] * temp[i]

    lower[n] = -q_t
    diag[n] = 1.0 + p_t + q_t
    rhs[n] = temp[n] + p_t * t_inf
    return thomas(lower, diag, upper, rhs)


def step_moisture(conc, c_inf, dt, hx, r, rf, picard_max, picard_tol):
    '''后向 Euler + Picard 迭代推进一步水分场（D 依赖 C）:
    1.结构类似于 step_temperature，采用后向 Euler 得到离散的水分方程
    2.D(C) 非线性，采用 Picard 迭代线性化，再使用 thomas 求解三对角方程组
    '''
    n = len(r) - 1
    g_c = dt / (hx * (RADIUS - hx / 4.0))
    p_c = 2.0 * g_c * RADIUS * H_MASS
    q_base = 2.0 * g_c * (RADIUS - hx / 2.0) / hx

    c_old = conc.copy()
    c_new = conc.copy()
    for _ in range(picard_max):
        d_node = diffusion_coefficient(c_new)
        d_face = 0.5 * (d_node[:-1] + d_node[1:])
        dtb = dt * d_face / (hx * hx)
        b0 = dt * d_face[0] / (hx * hx)

        lower = np.zeros(n + 1)
        diag = np.zeros(n + 1)
        upper = np.zeros(n + 1)
        rhs = np.zeros(n + 1)

        diag[0] = 1.0 + 4.0 * b0
        upper[0] = -4.0 * b0
        rhs[0] = c_old[0]
        for i in range(1, n):
            bm = dtb[i - 1] * rf[i - 1]
            bp = dtb[i] * rf[i]
            lower[i] = -bm
            upper[i] = -bp
            diag[i] = r[i] + bm + bp
            rhs[i] = r[i] * c_old[i]

        q_c = q_base * d_face[n - 1]
        lower[n] = -q_c
        diag[n] = 1.0 + p_c + q_c
        rhs[n] = c_old[n] + p_c * c_inf

        solved = thomas(lower, diag, upper, rhs)
        if np.max(np.abs(solved - c_new)) < picard_tol:
            c_new = solved
            break
        c_new = solved
    return c_new


def solve(n_fine=N_FINE, dt=DT, tend=T_END, picard_max=8, picard_tol=1e-13):
    #问题一求解
    hx = RADIUS / n_fine
    r = np.arange(n_fine + 1) * hx
    rf = (np.arange(n_fine) + 0.5) * hx

    stride = n_fine // N_OUT
    if stride * N_OUT != n_fine:
        raise ValueError("n_fine 必须是 N_OUT=%d 的整数倍" % N_OUT)

    t_air, temp_air, humid_air = load_air_conditions()

    n_steps = int(round(tend / dt))
    temp = np.full(n_fine + 1, T_INIT)
    conc = np.full(n_fine + 1, C_INIT)

    times = np.arange(1, n_steps + 1) * dt
    temp_out = np.empty((n_steps + 1, N_OUT + 1))
    conc_out = np.empty((n_steps + 1, N_OUT + 1))
    surf_flux = np.zeros(n_steps + 1)
    temp_out[0] = temp[::stride]
    conc_out[0] = conc[::stride]

    # 守恒检验用的细网格标量
    vol_factor = 2.0 * np.pi * LENGTH
    weights = cell_weights(n_fine, hx)
    c_volume = np.zeros(n_steps + 1)      # Q  = ∫ C dV
    phys_volume = np.zeros(n_steps + 1)   # W  = ∫ rho_d C dV
    c_volume[0] = vol_factor * float(np.sum(weights * conc))
    phys_volume[0] = vol_factor * RHO * float(
        np.sum(weights * conc / (1.0 + conc))
    )

    for step in range(1, n_steps + 1):
        t_now = step * dt
        t_inf = float(np.interp(t_now, t_air, temp_air))
        c_inf = float(np.interp(t_now, t_air, humid_air))
        conc = step_moisture(
            conc, c_inf, dt, hx, r, rf, picard_max, picard_tol
        )
        temp = step_temperature(temp, t_inf, dt, hx, r, rf)
        temp_out[step] = temp[::stride]
        conc_out[step] = conc[::stride]
        surf_flux[step] = H_MASS * (conc[-1] - c_inf)
        c_volume[step] = vol_factor * float(np.sum(weights * conc))
        phys_volume[step] = vol_factor * RHO * float(
            np.sum(weights * conc / (1.0 + conc))
        )

    return {
        "times": times,
        "r_out_cm": np.arange(N_OUT + 1) * (RADIUS / N_OUT) * 100.0,
        "temp_out": temp_out,
        "conc_out": conc_out,
        "surf_flux": surf_flux,
        "c_volume": c_volume,
        "phys_volume": phys_volume,
        "hx": hx,
        "dt": dt,
        "final_temp": temp,
        "final_conc": conc,
    }


def conservation_report(res):
    """水分守恒检验（在细网格上精确积分）。

    简单扩散形式的守恒量是 Q = ∫ C dV，其减少量应等于表面累计流出
        A * ∫ hm * (C_s - C_inf) dt,   A = 2*pi*R*L

    另报告物理水量 W = ∫ rho_d C dV = ∫ rho*C/(1+C) dV，用于说明
    简单模型与真实水分质量的差别（后者不守恒是密度口径近似的后果）。
    """
    area = 2.0 * np.pi * RADIUS * LENGTH
    q = res["c_volume"]
    removed = float(q[0] - q[-1])
    # 离散格式的守恒恒等式是左矩形和（通量取新时刻），
    # 用梯形法反而会引入 O(dt) 的假残差。
    cumulative = area * res["dt"] * float(np.sum(res["surf_flux"][1:]))
    phys = res["phys_volume"]
    return {
        "q_start": float(q[0]),
        "q_end": float(q[-1]),
        "removed": removed,
        "cumulative_flux": cumulative,
        "relative_residual": (removed - cumulative) / removed,
        "phys_start": float(phys[0]),
        "phys_end": float(phys[-1]),
    }


def print_report(res):
    #打印表 1 与表 2
    times = res["times"]
    pos_idx = [int(round(p / 0.1)) for p in REPORT_POS_CM]
    time_idx = [int(round(x / res["dt"])) for x in REPORT_TIMES]
    blocks = (
        ("表1  30 分钟内药材的温度", res["temp_out"], "°C"),
        ("表2  30 分钟内药材的水分浓度", res["conc_out"], "kg/kg"),
    )
    for title, hist, unit in blocks:
        print("\n=== %s  (单位: %s) ===" % (title, unit))
        header = "时间/s |" + "".join(
            "%10s" % ("%gcm" % p) for p in REPORT_POS_CM
        )
        print(header)
        print("-" * len(header))
        for ti in time_idx:
            row = "%6d |" % times[ti - 1]
            row += "".join("%10.4f" % hist[ti, pj] for pj in pos_idx)
            print(row)


def markdown_tables(res):
    #把表 1、表 2 渲染成 Markdown
    pos_idx = [int(round(p / 0.1)) for p in REPORT_POS_CM]
    time_idx = [int(round(x / res["dt"])) for x in REPORT_TIMES]
    out = [
        "# 问题一结果表",
        "",
        "由 `代码/solve_p1.py` 计算，细网格 $N=400$、$\\Delta t=1$ s，保留四位小数。",
        "",
    ]
    blocks = (
        ("表 1　30 分钟内药材的温度（单位：°C）", res["temp_out"]),
        ("表 2　30 分钟内药材的水分浓度（单位：kg/kg）", res["conc_out"]),
    )
    for title, hist in blocks:
        out.append("## " + title)
        out.append("")
        out.append(
            "| 时间/s | "
            + " | ".join("%g" % p for p in REPORT_POS_CM)
            + " |"
        )
        out.append("| --- |" + " --- |" * len(REPORT_POS_CM))
        for ti in time_idx:
            out.append(
                "| %d | " % res["times"][ti - 1]
                + " | ".join("%.4f" % hist[ti, pj] for pj in pos_idx)
                + " |"
            )
        out.append("")
    out.append("> 到药材中心的距离，单位 cm。完整结果见 `result1.xlsx`。")
    out.append("")
    return "\n".join(out)


def main():
    parser = argparse.ArgumentParser(description="A 题问题一求解")
    parser.add_argument("--n", type=int, default=N_FINE, help="细网格单元数")
    parser.add_argument("--dt", type=float, default=DT, help="时间步长 s")
    parser.add_argument("--save", default=None, help="保存解的 .npz 路径")
    parser.add_argument("--tables", default=None, help="导出表 1/表 2 的 Markdown 路径")
    args = parser.parse_args()

    print("求解中: N=%d, dt=%g s, 时长 %g s" % (args.n, args.dt, T_END))
    res = solve(n_fine=args.n, dt=args.dt)
    print_report(res)

    rep = conservation_report(res)
    print("\n=== 水分守恒检验（细网格精确积分）===")
    print("  模型守恒量 Q = ∫C dV   [m^3*kg/kg]")
    print("    初始        : %.8e" % rep["q_start"])
    print("    末态        : %.8e" % rep["q_end"])
    print("    减少        : %.8e" % rep["removed"])
    print("    表面累计流出: %.8e" % rep["cumulative_flux"])
    print("    相对残差    : %.3e" % rep["relative_residual"])
    print("  物理水量 W = ∫rho_d*C dV")
    print("    初始        : %.6f kg  (干料 %.6f kg)"
          % (rep["phys_start"], rep["phys_start"] / C_INIT))
    print("    末态        : %.6f kg" % rep["phys_end"])
    print("    脱除        : %.6f kg" % (rep["phys_start"] - rep["phys_end"]))

    print(
        "\n度量: alpha=%.4e m^2/s, D(C0)=%.4e m^2/s, Bi=%.3f, Bi_m=%.3f"
        % (
            ALPHA,
            diffusion_coefficient(C_INIT),
            H_HEAT * RADIUS / K_COND,
            H_MASS * RADIUS / diffusion_coefficient(C_INIT),
        )
    )

    if args.save:
        path = Path(args.save)
        path.parent.mkdir(parents=True, exist_ok=True)
        np.savez_compressed(
            path,
            times=res["times"],
            r_out_cm=res["r_out_cm"],
            temp_out=res["temp_out"],
            conc_out=res["conc_out"],
        )
        print("\n已保存: %s" % path)
        # 同时导出 JSON，供 代码/build_result1.py 生成 result1.xlsx
        import json

        json_path = path.with_suffix(".json")
        with json_path.open("w", encoding="utf-8") as fh:
            json.dump(
                {
                    "radius_cm": [round(float(v), 6) for v in res["r_out_cm"]],
                    "times": [int(v) for v in res["times"]],
                    # 去掉 t=0 的初始行，使第 i 行严格对应 times[i]
                    "temperature": np.round(res["temp_out"][1:], 4).tolist(),
                    "moisture": np.round(res["conc_out"][1:], 4).tolist(),
                },
                fh,
            )
        print("已保存: %s" % json_path)

    if args.tables:
        path = Path(args.tables)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(markdown_tables(res), encoding="utf-8")
        print("已保存: %s" % path)


if __name__ == "__main__":
    main()
